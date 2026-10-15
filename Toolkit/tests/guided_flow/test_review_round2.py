"""
Regression tests for the second review round: the Step 3 payload and counts, the Step 4
checks, the Step 1 SQL Server / Oracle pickers and File Library, and big-file support.

Run from the app folder:
    .venv\\Scripts\\python.exe -m unittest discover -s tests\\guided_flow -p test_review_round2.py -v
"""
import io
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from unittest import mock
from urllib.parse import quote_plus

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import pandas as pd  # noqa: E402

import helpers  # noqa: E402,F401  (installs the fake Salesforce)
import fake_salesforce  # noqa: E402
import isolation  # noqa: E402
from ui_components import data_operations as ops  # noqa: E402
from ui_components import large_data  # noqa: E402
from ui_components import step1_data_loading as step1  # noqa: E402
from ui_components import step4_verification as sv  # noqa: E402


def _bare_state():
    import streamlit as st
    st.session_state.setdefault('processing_status', {})
    for key in ('flow_load_result', 'flow_validation', 'insert_matching_strategy', 'insert_match_fields',
                'insert_target_field'):
        st.session_state.pop(key, None)
    return st.session_state


class PayloadTests(unittest.TestCase):
    """What Step 3 sends to Salesforce."""

    def setUp(self):
        isolation.isolate_file_library()
        self.state = _bare_state()
        fake_salesforce.reset()

    def test_blank_cells_are_null_not_nan(self):
        df = pd.read_csv(io.StringIO("Name,Active__c,Amount\nAcme,,\nGlobex,true,5\n"), dtype=str)
        df['When'] = pd.to_datetime(['2024-01-02 03:04:05', None])
        records = ops.clean_dataframe_for_salesforce(df).to_dict('records')
        json.dumps(records, allow_nan=False)   # a NaN would fail the whole batch
        self.assertIsNone(records[0]['Active__c'], "a blank checkbox must not become truthy")
        self.assertEqual(records[0]['When'], '2024-01-02 03:04:05')
        self.assertIsNone(records[1]['When'])

    def test_match_key_helper_column_is_not_sent(self):
        df = pd.DataFrame({'Name': ['A', 'B'], '_match_key': ['A|1', 'B|2']})
        sf = fake_salesforce.FakeSalesforce(username='u')
        ops.load_data_to_salesforce(sf, df, 'Account', 'Insert', 200, 1, field_mappings={},
                                    skip_lookup_resolution=True)
        sent = [r for _obj, _op, batch in fake_salesforce.WRITES for r in batch]
        self.assertEqual(len(sent), 2)
        self.assertTrue(all('_match_key' not in r for r in sent), sent)

    def test_failed_bulk_update_counts_each_row_once(self):
        sf = fake_salesforce.FakeSalesforce(username='u')
        df = pd.DataFrame({'Name': ['Acme', 'Nobody'], 'External_Key__c': ['K1', 'K-unknown']})

        def boom(self, records, **kwargs):
            raise RuntimeError("InvalidBatch")

        with mock.patch.object(fake_salesforce._FakeBulkType, 'update', boom):
            ops.load_data_to_salesforce(sf, df, 'Account', 'Update', 200, 1, update_match_field='External_Key__c',
                                        update_matching_strategy='external_id', field_mappings={},
                                        skip_lookup_resolution=True)
        load = self.state['flow_load_result']
        self.assertEqual(load['success'] + load['failed'] + load['skipped'], load['total'], load)
        self.assertEqual((load['failed'], load['skipped']), (1, 1))

    def test_rows_dropped_by_lookup_resolution_are_counted(self):
        sf = fake_salesforce.FakeSalesforce(username='u')
        df = pd.DataFrame({'Name': ['A', 'B', 'C']})

        def resolve(sf_conn, frame, target, mappings, return_stats=True):
            return frame.drop(index=[1]), {}, {}

        with mock.patch.object(ops, 'resolve_lookup_fields_with_mapping', resolve):
            ops.load_data_to_salesforce(sf, df, 'Account', 'Insert', 200, 1, field_mappings={'Name': 'Name'})
        load = self.state['flow_load_result']
        self.assertEqual((load['success'], load['skipped'], load['total']), (2, 1, 3))
        reasons = [r['reason'] for r in load['details']['skipped_rows']]
        self.assertTrue(any('lookup' in r for r in reasons), reasons)
        self.assertEqual([r['row'] for r in load['details']['skipped_rows']], [1])

    def test_match_settings_follow_the_strategy(self):
        self.state['insert_matching_strategy'] = 'field_combination'
        self.state['insert_match_fields'] = ['FirstName', 'LastName']
        self.state['insert_target_field'] = 'Combo__c'
        self.assertEqual(ops._load_match_settings('Insert', None, None, None, None, None, None)['fields'],
                         ['FirstName', 'LastName'])
        self.assertEqual(ops._load_match_settings('Update', None, None, 'field_concatenation',
                                                  ['FirstName', 'LastName'], None, None)['fields'],
                         ['FirstName', 'LastName'])
        self.assertEqual(ops._load_match_settings('Upsert', 'Ext__c', None, None, None, 'external_id',
                                                  None)['fields'], ['Ext__c'])


class SqlErrorHintTests(unittest.TestCase):

    def test_most_specific_hint_wins(self):
        db = ('[42000] [Microsoft][ODBC Driver 17 for SQL Server][SQL Server]Cannot open database "SalesDB" '
              'requested by the login. The login failed. (4060) (SQLDriverConnect); [28000] [Microsoft][ODBC '
              "Driver 17 for SQL Server][SQL Server]Login failed for user 'sa'. (18456)")
        self.assertIn('database name is wrong', ops.sql_error_hint(db))
        cert = ('[08001] [Microsoft][ODBC Driver 18 for SQL Server]SSL Provider: The certificate chain was '
                'issued by an authority that is not trusted.\r\n (-2146893019) (SQLDriverConnect)')
        self.assertIn('certificate was not accepted', ops.sql_error_hint(cert))
        login = "[28000] [Microsoft][ODBC Driver 17 for SQL Server][SQL Server]Login failed for user 'x'. (18456)"
        self.assertTrue(ops.sql_error_hint(login).startswith("Login failed for user 'x'"))


class _FakeSF:
    """A strict in-memory org for verify_salesforce_load (rejects duplicate SELECT fields like Salesforce)."""

    def __init__(self, records, fields):
        self.records = records
        self.fields = fields
        self.queries = []

    def __getattr__(self, name):
        if name.startswith('__'):
            raise AttributeError(name)
        outer = self

        class _Type:
            def describe(self):
                return {'fields': outer.fields}
        return _Type()

    def query_all(self, soql):
        self.queries.append(soql)
        select = soql.split('SELECT ', 1)[1].split(' FROM ', 1)[0]
        names = [f.strip().lower() for f in select.split(',')]
        if len(names) != len(set(names)):
            raise RuntimeError(f"MALFORMED_QUERY: duplicate field selected: {select}")
        where = soql.split(' WHERE ', 1)[1]
        field, values = where.split(' IN (', 1)
        wanted = {v.strip().strip("'") for v in values.rstrip(')').split(',')}
        rows = [r for r in self.records if str(r.get(field.strip())) in wanted]
        return {'records': [dict(r) for r in rows]}


def _field(name, ftype='string', **extra):
    return {'name': name, 'type': ftype, 'label': name, 'filterable': True, **extra}


class Step4CheckTests(unittest.TestCase):

    def _load(self, expected, ids, created=None, match=None, minutes_ago=1):
        started = datetime.now(timezone.utc) - timedelta(minutes=minutes_ago)
        return {'target': 'Account', 'operation': 'Insert', 'success': len(ids), 'failed': 0, 'skipped': 0,
                'total': len(ids), 'details': {
                    'record_ids': ids, 'expected': expected, 'row_positions': list(range(len(ids))),
                    'created': created or [True] * len(ids), 'match': match or {}, 'started_at_utc': started.isoformat()}}

    def test_audit_field_migration_neither_crashes_nor_fails(self):
        now = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%S.000+0000')
        record = {'Id': '001A', 'Name': 'Acme', 'CreatedDate': '2015-03-01T00:00:00.000+0000',
                  'LastModifiedDate': '2015-03-01T00:00:00.000+0000', 'SystemModstamp': now}
        sf = _FakeSF([record], [_field('Id', 'id'), _field('Name'), _field('CreatedDate', 'datetime'),
                                _field('LastModifiedDate', 'datetime'), _field('SystemModstamp', 'datetime')])
        expected = pd.DataFrame([{'Name': 'Acme', 'CreatedDate': '2015-03-01T00:00:00Z',
                                  'LastModifiedDate': '2015-03-01T00:00:00Z'}])
        result = sv.verify_salesforce_load(sf, self._load(expected, ['001A']), None)
        checks = {c['check']: c for c in result['checks']}
        self.assertEqual(checks["Field values match what was loaded"]['status'], 'PASS', checks)
        self.assertEqual(checks["Records were written by this load"]['status'], 'PASS', checks)

    def test_queries_stay_under_the_url_limit(self):
        values = [f"Company name with spaces, commas & quotes ' number {i}" for i in range(400)]
        records = [{'Id': f'001{i:015d}', 'Name': v, 'External_Key__c': v} for i, v in enumerate(values)]
        sf = _FakeSF(records, [_field('Id', 'id'), _field('Name'), _field('External_Key__c')])
        expected = pd.DataFrame([{'Name': v, 'External_Key__c': v} for v in values])
        load = self._load(expected, [r['Id'] for r in records], match={'fields': ['External_Key__c']})
        sv.verify_salesforce_load(sf, load, None)
        self.assertTrue(sf.queries)
        self.assertLessEqual(max(len(quote_plus(q)) for q in sf.queries), sv.MAX_SOQL_CHARS)

    def test_numbers_compare_within_the_field_precision(self):
        self.assertTrue(sv.values_match('10.005', 10.01, 'currency', 2))
        self.assertTrue(sv.values_match('10.004', 10.0, 'currency', 2))
        self.assertFalse(sv.values_match('10.02', 10.0, 'currency', 2))
        self.assertNotEqual(sv.normalize('1.2', 'double'), sv.normalize('1.4', 'double'),
                            "distinct numbers must not look like duplicate match keys")

    def test_sql_values_compare_by_column_type(self):
        canon = sv._sql_canon
        self.assertEqual(canon('10', 'decimal', 2), canon(Decimal('10.00'), 'decimal', 2))
        self.assertEqual(canon('1', 'bit'), canon(True, 'bit'))
        self.assertEqual(canon('2024-01-15 10:30:00.001', 'datetime'),
                         canon(datetime(2024, 1, 15, 10, 30), 'datetime'))
        self.assertNotEqual(canon('00123', 'nvarchar'), canon('123', 'nvarchar'), "leading zeros are data")
        expected = pd.DataFrame({'Amount': ['10', '2.5'], 'Code': ['007', '008']})
        table = pd.DataFrame({'Amount': [Decimal('10.00'), Decimal('2.50')], 'Code': ['007', '008']})

        def run_query(query):
            if 'INFORMATION_SCHEMA.TABLES' in query:
                return pd.DataFrame({'n': [1]})
            if 'INFORMATION_SCHEMA.COLUMNS' in query:
                return pd.DataFrame({'COLUMN_NAME': ['Amount', 'Code'], 'DATA_TYPE': ['decimal', 'nvarchar'],
                                     'NUMERIC_SCALE': [2, None]})
            if 'COUNT(*)' in query:
                return pd.DataFrame({'n': [2]})
            return table.copy()

        load = {'success': 2, 'failed': 0, 'skipped': 0, 'total': 2,
                'details': {'schema': 'dbo', 'table': 'T', 'mode': 'Append to Existing', 'before_count': 0,
                            'expected': expected}}
        result = sv.verify_sql_load(load, None, run_query)
        held = {c['check']: c for c in result['checks']}["Table holds the loaded rows"]
        self.assertEqual(held['status'], 'PASS', held)

    def test_sql_tests_never_switch_to_another_connection(self):
        with mock.patch('ui_components.connection_manager.sql_connections', return_value={}):
            with self.assertRaises(RuntimeError):
                sv._sql_runner({'details': {'sql_key': 'sql_Gone'}})

    def test_report_survives_control_characters(self):
        result = {'checks': [{'check': 'c', 'status': 'FAIL', 'detail': 'bad\x01value'}],
                  'details': {'mismatches': pd.DataFrame({'Sent': ['a\x0bb']})}}
        self.assertTrue(sv.report_bytes(result, [('Step 1 · source rows', 1)]).startswith(b'PK'))


class Step1SourceTests(unittest.TestCase):

    def test_select_detection_allows_leading_comments(self):
        self.assertTrue(step1._is_select_query("/****** Script for SelectTopNRows command from SSMS ******/\n"
                                               "SELECT TOP (1000) [Id] FROM [dbo].[T]"))
        self.assertTrue(step1._is_select_query("-- note\n  with x as (select 1) select * from x"))
        self.assertFalse(step1._is_select_query("DELETE FROM x"))

    def test_names_are_quoted_exactly(self):
        self.assertEqual(step1._bracket_parts(('dbo', 'a.b]c')), '[dbo].[a.b]]c]')
        self.assertEqual(step1._quote_oracle(('HR', 'Order "X"')), '"HR"."Order ""X"""')

    def test_oracle_list_logs_in_once_on_a_wrong_password(self):
        import streamlit as st
        st.session_state.clear() if hasattr(st.session_state, 'clear') else None
        calls = []

        def query(cfg, sql):
            calls.append(sql)
            raise RuntimeError("ORA-01017: invalid username/password; logon denied")

        with mock.patch.object(ops, '_oracle_query_to_df', query):
            tables, views, error, parts = step1._oracle_tables({'host': 'h', 'username': 'u'}, 'oracle_X',
                                                               force_refresh=True)
        self.assertEqual(len(calls), 1, "the pre-12c fallback must not cost a second failed login")
        self.assertIn('ORA-01017', error)

    def test_oracle_falls_back_before_12c(self):
        calls = []

        def query(cfg, sql):
            calls.append(sql)
            if 'oracle_maintained' in sql:
                raise RuntimeError('ORA-00904: "ORACLE_MAINTAINED": invalid identifier')
            return pd.DataFrame({'OWNER': ['HR'], 'OBJECT_NAME': ['EMP'], 'OBJECT_TYPE': ['TABLE']})

        with mock.patch.object(ops, '_oracle_query_to_df', query):
            tables, views, error, parts = step1._oracle_tables({'host': 'h2', 'username': 'u'}, 'oracle_Y',
                                                               force_refresh=True)
        self.assertEqual((len(calls), tables, error), (2, ['HR.EMP'], None))

    def test_file_library_reads_txt_like_an_upload(self):
        df = step1._parse_bytes(b"Name;City\nAcme;Pune\n", "accounts.txt", None)
        self.assertEqual(list(df.columns), ['Name', 'City'])


class BigFileTests(unittest.TestCase):

    def test_streamlit_lazy_tables_are_not_capped(self):
        from streamlit.delta_generator import DeltaGenerator
        self.assertTrue(large_data.streamlit_lazy_tables())
        large_data.install_dataframe_cap()
        self.assertFalse(getattr(DeltaGenerator.dataframe, '_dm_capped', False),
                         "capping would turn Streamlit's paged delivery into one big message")

    def test_excel_download_splits_over_sheets(self):
        from openpyxl import load_workbook
        from ui_components.data_hub import data_hub_ui
        with mock.patch.object(data_hub_ui, 'EXCEL_MAX_ROWS', 2):
            data = data_hub_ui._excel_bytes(pd.DataFrame({'a': range(5)}))
        self.assertEqual(load_workbook(io.BytesIO(data)).sheetnames, ['Data 1', 'Data 2', 'Data 3'])

    def test_datafiles_scan_indexes_txt_and_json(self):
        from ui_components import file_store_manager as fsm
        isolation.isolate_file_library()
        root = tempfile.mkdtemp(prefix='dm_datafiles_')
        os.makedirs(os.path.join(root, 'DataFiles', 'OrgA'))
        for name in ('a.txt', 'b.json', 'c.csv'):
            with open(os.path.join(root, 'DataFiles', 'OrgA', name), 'w') as handle:
                handle.write('x\n1\n')
        with mock.patch.object(fsm, 'PROJECT_ROOT', root):
            fsm.scan_and_register_existing()
        names = {f['original_name'] for f in fsm.list_source_files(org_name='OrgA')}
        self.assertTrue({'a.txt', 'b.json', 'c.csv'} <= names, names)


if __name__ == '__main__':
    unittest.main()
